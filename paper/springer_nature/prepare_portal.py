"""Create source-derived portal text and an editable cover-letter draft."""
from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT.parent / "reusability_report.tex"
PORTAL = ROOT / "portal"
LETTER = ROOT / "cover_letter"


def argument(text: str, command: str) -> str:
    match = re.search(re.escape(command) + r"\{", text)
    if match is None:
        raise ValueError(command)
    start = index = match.end()
    depth = 1
    while depth:
        if text[index] == "{" and text[index - 1] != "\\":
            depth += 1
        elif text[index] == "}" and text[index - 1] != "\\":
            depth -= 1
        index += 1
    return text[start:index - 1]


def plain(text: str) -> str:
    text = re.sub(r"\\href\{([^{}]*)\}\{([^{}]*)\}", r"\2: \1", text)
    text = text.replace(r"\noindent", "").replace(r"\%", "%")
    text = text.replace(r"\_", "_").replace(r"\&", "&")
    for command in ("path", "url", "texttt", "textbf", "textit"):
        text = re.sub(r"\\" + command + r"\{([^{}]*)\}", r"\1", text)
    text = text.replace("{,}", ",").replace("$", "")
    text = text.replace(chr(96) * 2, '"').replace("''", '"')
    text = text.replace("~", " ").replace("--", "-")
    if re.search(r"\\[A-Za-z]+", text):
        raise ValueError("Unconverted TeX in portal text: " + text)
    return re.sub(r"\s+", " ", text).strip()


def section(text: str, name: str) -> str:
    content = text.split("\\section*{" + name + "}", 1)[1]
    return plain(re.split(r"\\section\*|\\bibliography", content, maxsplit=1)[0])


def main() -> None:
    source = SOURCE.read_text()
    title = plain(argument(source, r"\title"))
    keywords = plain(argument(source, r"\keywords"))
    abstract = plain(source.split(r"\begin{abstract}", 1)[1].split(r"\end{abstract}", 1)[0])
    count = len(abstract.split())
    PORTAL.mkdir(exist_ok=True)
    for name, text in (("title.txt", title), ("abstract.txt", abstract), ("keywords.txt", keywords)):
        (PORTAL / name).write_text(text + "\n")
    fields = [
        ("TARGET JOURNAL", "Nature Machine Intelligence"),
        ("PROPOSED ARTICLE TYPE", "Reusability Report"),
        ("TITLE", title), (f"ABSTRACT ({count} whitespace-delimited words)", abstract),
        ("KEYWORDS", keywords),
        ("AUTHORS IN MANUSCRIPT ORDER",
         "1. Vatsal Pravinbhai Patel (corresponding author)\n2. Satya Prakash Mohanty\n"
         "3. Bikram Bhakat [2]\n4. AbhijeetKumar Pravinkumar Patel [1,3]\n"
         "5. Yash Jayeshbhai Patel [1,4]\n6. Ananya Pal [1,5]\n"
         "Vatsal Pravinbhai Patel and Satya Prakash Mohanty: affiliation [1]\n"
         "[1] HawkFranklin Research, India\n"
         "[2] Department of Zoology, The University of Burdwan, Purba Bardhaman, West Bengal, 713104, India\n"
         "[3] Primary Health Centre Dehri, Valsad, Gujarat, 396170, India\n"
         "[4] Medical College Baroda and SSG Hospital, Vadodara, Gujarat, 390001, India\n"
         "[5] Department of Zoology, Bethune College, Kolkata, West Bengal, 700006, India"),
        ("CORRESPONDING AUTHOR EMAIL", "vatsal1@hawkfranklin.in"),
        ("DATA AVAILABILITY", section(source, "Data availability")),
        ("CODE AVAILABILITY", section(source, "Code availability")),
        ("AUTHOR CONTRIBUTIONS", section(source, "Author contributions")),
        ("COMPETING INTERESTS", section(source, "Competing interests")),
    ]
    content = ("Submission portal copy from the current reusability report.\n"
               "Paste field text only; headings and instructions are not field content.\n"
               "The abstract is unchanged apart from conversion to plain text.\n\n")
    content += "\n\n".join(f"== {heading} ==\n{text}" for heading, text in fields) + "\n"
    (PORTAL / "portal_copy.txt").write_text(content)
    (PORTAL / "source_manifest.json").write_text(json.dumps({
        "source": str(SOURCE), "source_sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
        "abstract_word_count_whitespace": count,
        "policy": "No scientific editing; TeX converted to plain text only",
    }, indent=2) + "\n")
    # The example contains unrelated signed agreements. Reuse only its plain
    # letter pattern, never its signatures, revision claims or embedded images.
    document = Document()
    for s in document.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(1.0)
    style = document.styles["Normal"]
    # Arial: present in Word, and LibreOffice maps it to metric-identical Liberation Sans, so the PDF matches the .docx
    style.font.name, style.font.size = "Arial", Pt(10)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.08
    letter, _, authors = (LETTER / "cover_letter_portal.txt").read_text().partition("[[NEW PAGE]]")
    blocks = letter.strip().split("\n\n")
    for i, block in enumerate(blocks):
        p = document.add_paragraph()
        if i == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif block.startswith("Re: "):
            p.paragraph_format.keep_with_next = True
        if i >= len(blocks) - 3:                       # closing, sign-off and signature stay on one page
            p.paragraph_format.keep_with_next = i < len(blocks) - 1
        if i == len(blocks) - 1:
            p.paragraph_format.space_before = Pt(14)
        for j, line in enumerate(block.splitlines()):
            if j:
                p.add_run().add_break()
            run = p.add_run(line)
            if block.startswith("Re: "):
                run.bold = True
    # Second page: author list in manuscript order with contributions ("n | name | affiliations | contribution" rows -> table)
    for k, block in enumerate(authors.strip().split("\n\n") if authors.strip() else []):
        if k == 0:
            p = document.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.add_run(block).bold = True
            p.runs[0].font.size = Pt(12)
        elif " | " in block:
            rows = [line.split(" | ") for line in block.splitlines()]
            table = document.add_table(rows=len(rows) + 1, cols=4, style="Table Grid")
            table.autofit = False
            widths = (0.3, 1.7, 2.1, 2.4)
            grid = table._tbl.tblGrid.findall(qn("w:gridCol"))           # LibreOffice reads the grid, Word the cell widths
            for col, width in zip(grid, widths):
                col.set(qn("w:w"), str(int(width * 1440)))
            for c, (text, width) in enumerate(zip(("#", "Author", "Affiliation", "Contribution"), widths)):
                for r, values in enumerate([[text]] + [[row[c]] for row in rows]):
                    cell = table.cell(r, c)
                    cell.width = Inches(width)
                    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                    run = cell.paragraphs[0].add_run(values[0])
                    run.font.size, run.bold = Pt(9), r == 0
            document.add_paragraph().paragraph_format.space_after = Pt(0)
        else:
            document.add_paragraph(block)
    document.core_properties.title = "Cover letter for the cancer multiomics reusability report"
    document.core_properties.author = "Vatsal Pravinbhai Patel"
    document.core_properties.subject = "Draft submission to Nature Machine Intelligence"
    document.save(LETTER / "cover_letter.docx")
    print(f"Portal files created; unchanged abstract: {count} words. Cover letter drafted.")


if __name__ == "__main__":
    main()
